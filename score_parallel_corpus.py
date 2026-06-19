"""
韩中平行语料库质量评分程序

功能：
1. 对韩文句子进行质量评分
2. 对中文句子进行质量评分
3. 对韩中对齐质量进行评分
4. 计算最终总评分
5. 提供图形界面操作
"""

from __future__ import annotations
import math
import re
import unicodedata
from collections import Counter
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass
import threading

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("警告: openpyxl 未安装，将无法处理Excel文件")

# Konlpy 导入
try:
    from konlpy.tag import Okt
    KONLPY_AVAILABLE = True
except ImportError:
    KONLPY_AVAILABLE = False
    print("警告: Konlpy 不可用，韩文质量分析将使用基础方法")

# Jieba 导入
try:
    import jieba
    import jieba.posseg as pseg
    JIEBA_AVAILABLE = True
except ImportError:
    JIEBA_AVAILABLE = False
    print("警告: jieba 未安装，中文词汇分析将使用基础方法")

# GUI 相关
try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext
    HAS_GUI = True
except ImportError:
    HAS_GUI = False


# ===========================================================================
# 韩文评分模块
# ===========================================================================

_HANGUL_RE = re.compile(r'[\uAC00-\uD7A3]')
_SENTENCE_END_CHARS = {'.', '!', '?', '。', '！', '？'}


def clamp(value: float, min_value: float = 0.0, max_value: float = 1.0) -> float:
    return max(min_value, min(max_value, value))


def round2(value: float) -> float:
    return round(float(value), 2)


def saturation(value: float, target: float) -> float:
    if target <= 0:
        return 0.0
    return clamp(value / target)


def category_scope(category: int) -> float:
    if category <= 1:
        return 0.35
    if category == 2:
        return 0.70
    return 1.0



def anomaly_length_score(text_len: int, max_len: int, max_score: float = 5.0) -> float:
    """长度只做异常检测，不再承担主要质量判断。"""
    if text_len <= 0:
        return 0.0
    if text_len <= 2:
        return max_score * 0.45
    if text_len <= 4:
        return max_score * 0.72
    if text_len <= max_len:
        return max_score
    overflow_ratio = (text_len - max_len) / max_len
    return max_score * clamp(1.0 - overflow_ratio * 2.5)


def centered_ratio_score(ratio: float, max_score: float, sharpness: float = 1.35) -> float:
    """以 1 为中心的对称长度比评分。"""
    if ratio <= 0:
        return 0.0
    return max_score * math.exp(-sharpness * abs(math.log(ratio)))


def multiset_similarity(seq1, seq2) -> float:
    if not seq1 and not seq2:
        return 1.0
    counter1 = Counter(seq1)
    counter2 = Counter(seq2)
    keys = set(counter1) | set(counter2)
    numerator = sum(min(counter1[k], counter2[k]) for k in keys)
    denominator = sum(max(counter1[k], counter2[k]) for k in keys)
    return numerator / denominator if denominator else 1.0


def extract_normalized_punctuation(text: str):
    normalized = unicodedata.normalize("NFKC", text or "")
    return [ch for ch in normalized if unicodedata.category(ch).startswith('P')]


def normalize_inner_punctuation(ch: str) -> str:
    mapping = {
        '，': ',',
        ',': ',',
        '、': ',',
        '：': ':',
        ':': ':',
        '；': ';',
        ';': ';',
        '（': '(',
        '）': ')',
        '【': '[',
        '】': ']',
        '「': '"',
        '」': '"',
        '『': '"',
        '』': '"',
        '“': '"',
        '”': '"',
        '‘': "'",
        '’': "'",
        '《': '"',
        '》': '"',
    }
    return mapping.get(ch, ch)


def normalize_end_punctuation(ch: str) -> str:
    mapping = {
        '.': '.',
        '。': '.',
        '．': '.',
        '?': '?',
        '？': '?',
        '!': '!',
        '！': '!',
    }
    return mapping.get(ch, "")


def last_end_punctuation(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text or "").strip()
    trailing_quotes = {'"', "'", '”', '’', '」', '』', '》', ')', '）', ']', '】'}
    for ch in reversed(normalized):
        if ch in trailing_quotes:
            continue
        normalized_ch = normalize_end_punctuation(ch)
        if normalized_ch:
            return normalized_ch
    return ""


def extract_inner_punctuation(text: str):
    normalized = unicodedata.normalize("NFKC", text or "").strip()
    trailing_quotes = {'"', "'", '”', '’', '」', '』', '》', ')', '）', ']', '】'}
    end_idx = -1

    i = len(normalized) - 1
    while i >= 0 and normalized[i] in trailing_quotes:
        i -= 1
    if i >= 0 and normalize_end_punctuation(normalized[i]):
        end_idx = i

    result = []
    for idx, ch in enumerate(normalized):
        if not unicodedata.category(ch).startswith('P'):
            continue
        # 排除句末真正的结束标点，以及其后仅用于收尾的引号/括号
        if end_idx != -1 and idx >= end_idx:
            continue
        normalized_ch = normalize_inner_punctuation(ch)
        if normalize_end_punctuation(normalized_ch):
            continue
        # 内部引号在韩中互译中经常一边保留、一边省略，不作为强匹配条件
        if normalized_ch in {'"', "'"}:
            continue
        result.append(normalized_ch)
    return result


def inner_punctuation_similarity(seq1, seq2) -> float:
    raw = multiset_similarity(seq1, seq2)
    if raw > 0:
        return raw
    if not seq1 and not seq2:
        return 1.0

    minor_marks = {',', ':', ';', '"', "'", '(', ')', '[', ']'}
    merged = list(seq1) + list(seq2)
    if merged and all(ch in minor_marks for ch in merged):
        length_gap = abs(len(seq1) - len(seq2))
        if length_gap <= 1:
            return 0.70
        if length_gap <= 2:
            return 0.45
    return 0.0


def has_bilingual_note(text: str) -> bool:
    t = unicodedata.normalize("NFKC", text or "").strip()
    pattern = r'^[\w\s\u4e00-\u9fff]*\([\w\s\u4e00-\u9fff]+\)[\w\s\u4e00-\u9fff]*$'
    return bool(re.match(pattern, t))


class KoreanScorer:
    """韩文句子评分器"""
    
    def __init__(self):
        self.okt = None
        if KONLPY_AVAILABLE:
            try:
                self.okt = Okt()
                print("✅ Konlpy Okt初始化成功")
            except Exception as e:
                print(f"⚠️  Konlpy初始化失败: {e}")
        self.copula_words = [
            # 系动词/存在动词
            '이다', '있', '없', '되'
        ]
        self.action_verb_words = [
            # 动作动词词干
            '가', '오', '먹', '마시', '보', '듣', '말', '하',
            '쓰', '읽', '배우', '가르치', '묻', '답하',
            '사', '팔', '주', '받', '놓', '들', '밀', '당기',
            '열', '닫', '움직이', '만들', '바꾸'
        ]
        self.mental_verb_words = [
            # 心理/认知动词
            '생각', '좋아', '사랑', '알', '모르', '느끼', '믿',
            '이해', '원하', '바라', '필요', '기억', '잊', '걱정'
        ]
        self.adjective_words = [
            # 形容词词干
            '좋', '나쁘', '크', '작', '많', '적',
            '빠르', '느리', '높', '낮', '길', '짧',
            '예쁘', '아름답', '못생기', '행복', '슬프',
            '기쁘', '어렵', '쉽', '깨끗하', '더럽', '조용하'
        ]
        self.degree_adverb_words = [
            # 程度/否定/时态副词
            '아주', '매우', '너무', '정말', '참',
            '조금', '약간', '좀', '거의', '다', '모두',
            '안', '못', '말고', '이미', '지금', '곧'
        ]
        self.ending_particle_words = [
            # 语尾/助词/连接成分
            '다', '요', '아', '어', 'ㄴ다', '는다', 'ㄹ다',
            '었', '았', '였', '겠',
            '이', '가', '은', '는', '을', '를', '에', '에서',
            '와', '과', '랑', '하고', '의', '에게', '한테',
            '까지', '부터', '만', '도'
        ]
        self.predicate_words = (
            self.copula_words +
            self.action_verb_words +
            self.mental_verb_words +
            self.adjective_words +
            self.degree_adverb_words +
            self.ending_particle_words
        )
    
    def analyze_quality(self, text: str):
        """分析韩文句子质量 - 分词器优先，不足时基础词池补足"""
        result = {
            'has_noun': False,
            'has_verb': False,
            'has_adj': False,
            'noun_count': 0,
            'predicate_count': 0,
            'content_word_count': 0,
            'is_good_quality': False
        }

        text = text.strip()
        if not text:
            return result

        hangul_count = len(_HANGUL_RE.findall(text))
        if hangul_count == 0:
            return result

        if self.okt:
            try:
                pos_tags = self.okt.pos(text)
                for word, flag in pos_tags:
                    if flag in ['Noun', 'NNP', 'NNG']:
                        result['noun_count'] += 1
                        result['has_noun'] = True
                        result['content_word_count'] += 1
                    elif flag in ['Verb', 'VV']:
                        result['has_verb'] = True
                        result['predicate_count'] += 1
                        result['content_word_count'] += 1
                    elif flag in ['Adjective', 'VA']:
                        result['has_adj'] = True
                        result['predicate_count'] += 1
                        result['content_word_count'] += 1
                    elif flag in ['Adverb', 'Josa', 'Eomi', 'Exclamation']:
                        result['content_word_count'] += 1

                if result['noun_count'] == 0 or result['predicate_count'] == 0:
                    basic_hits = sum(1 for word in self.predicate_words if word in text)
                    result['predicate_count'] = max(result['predicate_count'], basic_hits)
                    if basic_hits > 0:
                        result['has_verb'] = True

                if result['noun_count'] == 0:
                    result['noun_count'] = max(1, hangul_count // 3)
                    result['has_noun'] = True
                if result['content_word_count'] == 0:
                    result['content_word_count'] = max(1, result['noun_count'] + result['predicate_count'])

                result['is_good_quality'] = (
                    result['has_noun'] and
                    (result['has_verb'] or result['has_adj'] or result['predicate_count'] >= 1)
                )
                return result
            except Exception:
                pass

        basic_hits = []
        for word in self.predicate_words:
            if word in text:
                basic_hits.append(word)

        result['predicate_count'] = len(set(basic_hits))
        result['has_verb'] = any(word in text for word in (self.copula_words + self.action_verb_words + self.mental_verb_words))
        result['has_adj'] = any(word in text for word in self.adjective_words)
        result['noun_count'] = max(1, hangul_count // 3)
        result['has_noun'] = result['noun_count'] > 0
        result['content_word_count'] = max(
            result['noun_count'],
            min(20, result['noun_count'] + result['predicate_count'])
        )
        result['is_good_quality'] = result['has_noun'] and (result['has_verb'] or result['has_adj'] or result['predicate_count'] >= 1)
        return result
    
    def score(self, text: str):
        """韩文句子评分 - 连续小数分，长度仅做异常检测"""
        details = []
        
        text = text.strip()
        text_len = len(text)
        
        # 基础统计
        hangul_count = len(_HANGUL_RE.findall(text))
        hangul_ratio = hangul_count / text_len if text_len > 0 else 0
        quality = self.analyze_quality(text)
        unique_hangul = len(set(_HANGUL_RE.findall(text)))
        has_predicate = quality['predicate_count'] >= 1 or quality['has_verb'] or quality['has_adj']
        has_sentence_end = text[-1:] in _SENTENCE_END_CHARS
        
        # 1. 长度分（5分）- 仅用于检测过短/超长异常
        score_length = anomaly_length_score(text_len, 128, 5.0)
        details.append(f"长度:{score_length:.2f}")
        
        # 2. 韩文比例分（25分）- 连续按韩文占比给分
        score_hangul = 25.0 * (clamp((hangul_ratio - 0.08) / 0.87) ** 0.9)
        details.append(f"韩文比例:{score_hangul:.2f}")
        
        # 3. 名词分（8分）- 仅体现丰富性
        noun_signal = quality['noun_count'] + min(unique_hangul, 12) * 0.12
        score_noun = 8.0 * saturation(noun_signal, 5.0)
        details.append(f"名词:{score_noun:.2f}")
        
        # 4. 动形分（8分）- 仅作辅助信号
        predicate_signal = (
            quality['predicate_count'] +
            0.7 * float(quality['has_verb']) +
            0.5 * float(quality['has_adj'])
        )
        score_verb_adj = 8.0 * saturation(predicate_signal, 4.0)
        details.append(f"动形:{score_verb_adj:.2f}")
        
        # 5. 内容词分（8分）- 作为丰富度参考
        content_signal = quality['content_word_count'] + min(unique_hangul, 15) * 0.10
        score_content = 8.0 * saturation(content_signal, 8.0)
        details.append(f"内容词:{score_content:.2f}")
        
        # 6. 完整性分（46分）- 核心质量项
        completeness_signal = (
            0.28 * saturation(hangul_count, 6.0) +
            0.18 * hangul_ratio +
            0.18 * float(quality['has_noun']) +
            0.20 * float(has_predicate) +
            0.10 * float(has_sentence_end) +
            0.06 * saturation(quality['predicate_count'], 2.5)
        )
        score_completeness = 46.0 * clamp(completeness_signal)
        details.append(f"完整:{score_completeness:.2f}")

        score_total = round2(min(
            score_length +
            score_hangul +
            score_noun +
            score_verb_adj +
            score_content +
            score_completeness,
            100.0
        ))
        
        return {
            'total': score_total,
            'length': round2(score_length),
            'hangul_ratio': round2(score_hangul),
            'noun': round2(score_noun),
            'verb_adj': round2(score_verb_adj),
            'content_word': round2(score_content),
            'completeness': round2(score_completeness),
            'details': '|'.join(details)
        }


# ===========================================================================
# 中文评分模块
# ===========================================================================

_CHINESE_RE = re.compile(r'[\u4e00-\u9fff]')

DEFAULT_KO_COLUMN = 2
DEFAULT_ZH_COLUMN = 4
KO_COLUMN_CANDIDATES = ["韩文修正", "韩文", "KO修正", "KO", "원문", "한국어"]
ZH_COLUMN_CANDIDATES = ["ZH修正", "中文修正", "中文", "ZH", "译文", "翻译"]


def find_column_by_headers(ws, candidates, fallback):
    """优先按列名匹配，匹配不到时回退到固定列号。"""
    header_map = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value is None:
            continue
        header = str(value).strip()
        if header:
            header_map[header] = col

    for candidate in candidates:
        if candidate in header_map:
            return header_map[candidate], candidate, True

    return fallback, f"fallback:{fallback}", False


class ChineseScorer:
    """中文句子评分器"""

    def __init__(self):
        self.jieba_available = False
        if JIEBA_AVAILABLE:
            try:
                jieba.initialize()
                self.jieba_available = True
                print("✅ jieba初始化成功")
            except Exception as e:
                print(f"⚠️  jieba初始化失败: {e}，使用基础方法")

        self.copula_words = [
            # 系动词/存在动词
            '是', '为', '乃', '即', '有', '存在', '拥有', '具备',
        ]
        self.action_verb_words = [
            # 动作动词
            '在', '做', '说', '想', '看', '听', '走', '跑', '跳', '飞',
            '来', '去', '进', '出', '上', '下', '起', '落',
            '吃', '喝', '睡', '醒', '坐', '站', '躺', '写', '读', '学', '教',
            '买', '卖', '送', '给', '拿', '取', '放', '装',
            '开', '关', '打', '拉', '推', '搬', '运',
            '支持', '反对', '保持', '提供', '产生', '形成', '实现', '改善',
        ]
        self.mental_verb_words = [
            # 心理/认知动词
            '问', '答', '喜欢', '爱', '恨', '讨厌', '害怕', '担心',
            '知道', '明白', '了解', '理解', '懂得', '认识',
            '需要', '想要', '希望', '渴望', '盼望', '必须', '应该',
            '可以', '能够', '会', '能', '觉得', '感觉', '感到', '认为', '以为',
            '看见', '听到', '闻到', '尝到', '碰到',
        ]
        self.adjective_words = [
            # 形容词/状态词
            '好', '坏', '大', '小', '多', '少', '快', '慢', '高', '低',
            '长', '短', '宽', '窄', '厚', '薄', '深', '浅',
            '美', '丑', '漂亮', '难看', '可爱',
            '高兴', '开心', '快乐', '悲伤', '难过', '伤心',
            '热闹', '安静', '干净', '脏',
            '新', '旧', '老', '年轻', '新鲜',
            '热', '冷', '暖', '凉', '甜', '咸', '酸', '辣', '苦', '香',
            '亮', '暗', '明', '黑', '白',
        ]
        self.degree_adverb_words = [
            # 程度/否定/时态副词
            '很', '非常', '特别', '十分', '极其', '比较', '更', '最',
            '不', '没', '没有', '别', '不要',
            '都', '全', '已经', '正在', '将要', '即将',
        ]
        self.aspect_particle_words = [
            # 体貌/语气/补语
            '了', '着', '过', '起来', '下去', '下来', '上来',
            '吗', '呢', '吧', '啊', '呀'
        ]
        self.predicate_words = (
            self.copula_words +
            self.action_verb_words +
            self.mental_verb_words +
            self.adjective_words +
            self.degree_adverb_words +
            self.aspect_particle_words
        )

    def analyze_quality(self, text: str):
        """分析中文句子质量"""
        result = {
            'has_noun': False,
            'has_verb': False,
            'has_adj': False,
            'noun_count': 0,
            'predicate_count': 0,
            'content_word_count': 0,
            'is_good_quality': False
        }

        text = text.strip()
        if not text:
            return result

        chinese_count = len(_CHINESE_RE.findall(text))
        if chinese_count == 0:
            return result

        # 优先使用 jieba 的词性标注
        if self.jieba_available:
            try:
                pos_tags = list(pseg.cut(text))
                for word, flag in pos_tags:
                    if flag.startswith('n'):
                        result['noun_count'] += 1
                        result['has_noun'] = True
                        result['content_word_count'] += 1
                    elif flag.startswith('v'):
                        result['has_verb'] = True
                        result['predicate_count'] += 1
                        result['content_word_count'] += 1
                    elif flag.startswith('a'):
                        result['has_adj'] = True
                        result['predicate_count'] += 1
                        result['content_word_count'] += 1
                    elif flag in ['r', 'm', 'q', 't', 's', 'l', 'i']:
                        result['content_word_count'] += 1

                result['is_good_quality'] = (
                    result['noun_count'] >= 1 and
                    (result['has_verb'] or result['has_adj'] or result['predicate_count'] >= 1)
                )

                # jieba 给出的结果过少时，补基础检测
                if result['noun_count'] == 0 or result['predicate_count'] == 0:
                    basic_hits = sum(1 for word in self.predicate_words if word in text)
                    result['predicate_count'] = max(result['predicate_count'], basic_hits)
                    if basic_hits > 0:
                        result['has_verb'] = True

                if result['noun_count'] == 0:
                    result['noun_count'] = max(1, chinese_count // 4)
                    result['has_noun'] = True
                if result['content_word_count'] == 0:
                    result['content_word_count'] = max(1, chinese_count // 3)

                result['is_good_quality'] = (
                    result['has_noun'] and
                    (result['has_verb'] or result['has_adj'] or result['predicate_count'] >= 1)
                )
                return result
            except Exception:
                pass

        # jieba 不可用或失败时，使用基础谓词策略
        basic_hits = []
        for word in self.predicate_words:
            if word in text:
                basic_hits.append(word)

        result['predicate_count'] = len(set(basic_hits))
        result['has_verb'] = result['predicate_count'] > 0
        result['has_adj'] = any(word in text for word in self.adjective_words)
        result['noun_count'] = max(1, chinese_count // 4)
        result['has_noun'] = result['noun_count'] > 0
        result['content_word_count'] = max(
            result['noun_count'],
            min(20, chinese_count // 3 + result['predicate_count'])
        )
        result['is_good_quality'] = result['has_noun'] and (result['has_verb'] or result['has_adj'])
        return result

    def score(self, text: str):
        """中文句子评分 - 连续小数分，长度仅做异常检测"""
        details = []

        text = text.strip()
        text_len = len(text)

        chinese_count = len(_CHINESE_RE.findall(text))
        chinese_ratio = chinese_count / text_len if text_len > 0 else 0
        quality = self.analyze_quality(text)
        unique_chinese = len(set(_CHINESE_RE.findall(text)))
        has_sentence_end = text[-1:] in _SENTENCE_END_CHARS
        has_alnum = bool(re.search(r"[A-Za-z0-9]", text))

        # 1. 长度分（5分）- 仅用于检测过短/超长异常
        score_length = anomaly_length_score(text_len, 96, 5.0)
        details.append(f"长度:{score_length:.2f}")

        # 2. 中文比例分（25分）- 连续按中文占比给分
        score_chinese = 25.0 * (clamp((chinese_ratio - 0.08) / 0.87) ** 0.9)
        details.append(f"中文比例:{score_chinese:.2f}")

        # 3. 名词分（8分）- 只体现丰富度
        noun_signal = quality['noun_count'] + min(unique_chinese, 12) * 0.12
        score_noun = 8.0 * saturation(noun_signal, 5.0)
        details.append(f"名词:{score_noun:.2f}")

        # 4. 谓词分（8分）- 只作辅助
        predicate_signal = (
            quality['predicate_count'] +
            0.7 * float(quality['has_verb']) +
            0.5 * float(quality['has_adj'])
        )
        score_predicate = 8.0 * saturation(predicate_signal, 4.0)
        details.append(f"谓词:{score_predicate:.2f}")

        # 5. 内容词分（8分）- 作为丰富度参考
        content_signal = quality['content_word_count'] + min(unique_chinese, 15) * 0.10
        score_content = 8.0 * saturation(content_signal, 8.0)
        details.append(f"内容词:{score_content:.2f}")

        # 6. 完整性分（46分）- 核心质量项
        completeness_signal = (
            0.28 * saturation(chinese_count, 6.0) +
            0.18 * chinese_ratio +
            0.18 * float(quality['has_noun']) +
            0.20 * float(quality['has_verb'] or quality['has_adj'] or quality['predicate_count'] >= 1) +
            0.10 * float(has_sentence_end) +
            0.06 * saturation(quality['predicate_count'], 2.5)
        )
        structure_factor = clamp((chinese_count / 12.0) ** 0.7)
        if chinese_count < 6 and (has_alnum or chinese_ratio < 0.6):
            structure_factor *= 0.6
        score_completeness = 46.0 * clamp(completeness_signal) * clamp(structure_factor)
        details.append(f"完整:{score_completeness:.2f}")

        score_total = round2(min(
            score_length +
            score_chinese +
            score_noun +
            score_predicate +
            score_content +
            score_completeness,
            100.0
        ))

        return {
            'total': score_total,
            'length': round2(score_length),
            'chinese_ratio': round2(score_chinese),
            'noun': round2(score_noun),
            'predicate': round2(score_predicate),
            'content_word': round2(score_content),
            'completeness': round2(score_completeness),
            'details': '|'.join(details)
        }


# ===========================================================================
# 对齐质量评分模块
# ===========================================================================

class AlignmentScorer:
    """韩中对齐质量评分器"""
    
    def __init__(self):
        pass
    
    def extract_special_elements(self, text: str):
        """提取特殊元素"""
        numbers = re.findall(r'\d+', text)
        english = re.findall(r'[a-zA-Z]+', text)
        brackets = re.findall(r'[\[\](){}（）【】]', text)
        return {
            'numbers': numbers,
            'english': english,
            'brackets': brackets
        }
    
    def check_sentence_type(self, text: str):
        """判断句子类型"""
        t = unicodedata.normalize("NFKC", text or "")
        comma_count = t.count(",")
        has_strong_punc = any(ch in t for ch in [";", "；", ":", "："])

        ko_strong_connectives = ['아서', '어서', '니', '지만', '는데', '거나']
        cn_strong_connectives = ['因为', '所以', '虽然', '但是', '但', '却', '而']

        has_strong_connective = any(conn in t for conn in (ko_strong_connectives + cn_strong_connectives))

        if has_strong_punc:
            return 'complex'
        if comma_count >= 2:
            return 'complex'
        if has_strong_connective and comma_count >= 1:
            return 'complex'
        return 'simple'

    def sentence_complexity(self, text: str) -> float:
        t = unicodedata.normalize("NFKC", text or "")
        ko_strong_connectives = ['아서', '어서', '니', '지만', '는데', '거나']
        cn_strong_connectives = ['因为', '所以', '虽然', '但是', '但', '却', '而']

        score = 0.0
        comma_count = t.count(",")
        if comma_count >= 2:
            score += 0.35
        elif comma_count == 1:
            score += 0.15
        if ';' in t or '；' in t or ':' in t or '：' in t:
            score += 0.15

        conn_count = 0
        for conn in ko_strong_connectives + cn_strong_connectives:
            if conn and conn in t:
                conn_count += 1
        score += min(0.5, 0.18 * conn_count)

        if len(t.strip()) >= 40:
            score += 0.10
        return clamp(score)
    
    def has_question_mark(self, text: str):
        """检查是否有问号"""
        return '?' in text or '？' in text
    
    def check_english_chinese_pairs(self, ko_text: str, zh_text: str):
        """检查双语注释"""
        pattern = r'^[\w\s\u4e00-\u9fff]*\([\w\s\u4e00-\u9fff]+\)[\w\s\u4e00-\u9fff]*$'
        if re.match(pattern, ko_text) or re.match(pattern, zh_text):
            return True
        return False
    
    def score(self, ko_text: str, zh_text: str):
        """对齐质量评分 - 使用连续小数分"""
        details = []
        
        ko_text = ko_text or ""
        zh_text = zh_text or ""
        
        # 1. 长度匹配（5分）- 以 1 为中心的连续对称评分
        ko_len = len(ko_text.strip())
        zh_len = len(zh_text.strip())
        
        if ko_len > 0 and zh_len > 0:
            ratio = zh_len / ko_len
            score_length = centered_ratio_score(ratio, 5.0, 1.4)
        else:
            score_length = 0.0
        details.append(f"长度匹配:{score_length:.2f}")
        
        # 2. 特殊元素保留（10分）- 保留重要性，但继续降权
        ko_elems = self.extract_special_elements(ko_text)
        zh_elems = self.extract_special_elements(zh_text)
        num_sim = multiset_similarity(ko_elems['numbers'], zh_elems['numbers'])
        en_sim = multiset_similarity([s.lower() for s in ko_elems['english']], [s.lower() for s in zh_elems['english']])
        bracket_sim = multiset_similarity(ko_elems['brackets'], zh_elems['brackets'])
        score_special = 10.0 * (0.45 * num_sim + 0.35 * en_sim + 0.20 * bracket_sim)
        details.append(f"特殊元素:{score_special:.2f}")
        
        # 3. 标点符号匹配（10分）- 句末标点等价 + 句内标点一致
        ko_end = last_end_punctuation(ko_text)
        zh_end = last_end_punctuation(zh_text)
        ko_inner_punc = extract_inner_punctuation(ko_text)
        zh_inner_punc = extract_inner_punctuation(zh_text)
        inner_punc_sim = inner_punctuation_similarity(ko_inner_punc, zh_inner_punc)
        has_inner_punc = bool(ko_inner_punc or zh_inner_punc)

        if ko_end and zh_end and ko_end == zh_end:
            score_punc = 10.0
        elif ko_end and zh_end:
            if inner_punc_sim >= 0.15:
                score_punc = 9.0
            else:
                score_punc = 8.0
        elif ko_end or zh_end:
            if inner_punc_sim >= 0.20:
                score_punc = 9.0
            else:
                score_punc = 8.0
        else:
            score_punc = 8.0 if inner_punc_sim >= 0.20 else 7.0
        details.append(f"标点匹配:{score_punc:.2f}")
        
        # 4. 问号匹配（7分）- 保留但进一步降权
        ko_has_q = self.has_question_mark(ko_text)
        zh_has_q = self.has_question_mark(zh_text)

        if ko_has_q and zh_has_q:
            score_qmark = 7.0
        elif not ko_has_q and not zh_has_q:
            score_qmark = 7.0
        elif ko_has_q and not zh_has_q:
            score_qmark = 2.0
        else:
            score_qmark = 5.0
        details.append(f"问号匹配:{score_qmark:.2f}")
        
        # 5. 句子类型匹配（56分）
        ko_type = self.check_sentence_type(ko_text)
        zh_type = self.check_sentence_type(zh_text)
        ko_hangul = len(_HANGUL_RE.findall(ko_text))
        zh_chinese = len(_CHINESE_RE.findall(zh_text))
        zh_has_alnum = bool(re.search(r"[A-Za-z0-9]", zh_text))
        zh_frag = zh_chinese < 6 and (zh_has_alnum or (len(zh_text.strip()) > 0 and zh_chinese / max(1, len(zh_text.strip())) < 0.6))
        ko_frag = ko_hangul < 6 and len(ko_text.strip()) < 12

        if ko_type == zh_type:
            score_type = 56.0
        else:
            ko_c = self.sentence_complexity(ko_text)
            zh_c = self.sentence_complexity(zh_text)
            if abs(ko_c - zh_c) < 0.25:
                score_type = 42.0
            else:
                score_type = 28.0
        if zh_frag or ko_frag:
            score_type = min(score_type, 28.0)
        details.append(f"句型匹配:{score_type:.2f}")
        
        # 6. 双语注释宽容度（12分）- 提高宽容度权重
        ko_note = has_bilingual_note(ko_text)
        zh_note = has_bilingual_note(zh_text)
        if ko_note == zh_note:
            score_bilingual = 12.0
        else:
            score_bilingual = 10.0
        details.append(f"双语注释:{score_bilingual:.2f}")

        score_total = round2(min(
            score_length +
            score_special +
            score_punc +
            score_qmark +
            score_type +
            score_bilingual,
            100.0
        ))
        
        return {
            'total': score_total,
            'length_match': round2(score_length),
            'special_elements': round2(score_special),
            'punctuation_match': round2(score_punc),
            'question_mark_match': round2(score_qmark),
            'sentence_type_match': round2(score_type),
            'bilingual_note': round2(score_bilingual),
            'details': '|'.join(details)
        }


# ===========================================================================
# 主评分类
# ===========================================================================

@dataclass
class ParallelScoreResult:
    """平行语料评分结果"""
    ko_score_total: float
    ko_score_length: float
    ko_score_hangul: float
    ko_score_noun: float
    ko_score_verb_adj: float
    ko_score_content_word: float
    ko_score_completeness: float
    ko_details: str
    
    zh_score_total: float
    zh_score_length: float
    zh_score_chinese_ratio: float
    zh_score_noun: float
    zh_score_predicate: float
    zh_score_content_word: float
    zh_score_completeness: float
    zh_details: str
    
    align_score_total: float
    align_score_length: float
    align_score_special: float
    align_score_punctuation: float
    align_score_question_mark: float
    align_score_sentence_type: float
    align_score_bilingual: float
    align_details: str
    
    final_score: float
    final_details: str


class ParallelCorpusScorer:
    """平行语料库评分器"""
    
    def __init__(self):
        self.ko_scorer = KoreanScorer()
        self.zh_scorer = ChineseScorer()
        self.align_scorer = AlignmentScorer()
    
    def score(self, ko_text: str, zh_text: str):
        """对韩中平行句子进行评分 - 对齐权重最高"""
        
        ko_result = self.ko_scorer.score(ko_text or "")
        zh_result = self.zh_scorer.score(zh_text or "")
        align_result = self.align_scorer.score(ko_text or "", zh_text or "")
        
        # 对齐权重最高，适合模型微调
        final_score = round2(
            ko_result['total'] * 0.25 +
            zh_result['total'] * 0.3 +
            align_result['total'] * 0.45
        )
        
        final_details = (
            f"韩文:{ko_result['total']:.2f}分(25%)|"
            f"中文:{zh_result['total']:.2f}分(30%)|"
            f"对齐:{align_result['total']:.2f}分(45%)|"
            f"总分:{final_score:.2f}分"
        )
        
        return ParallelScoreResult(
            ko_score_total=ko_result['total'],
            ko_score_length=ko_result['length'],
            ko_score_hangul=ko_result['hangul_ratio'],
            ko_score_noun=ko_result['noun'],
            ko_score_verb_adj=ko_result['verb_adj'],
            ko_score_content_word=ko_result['content_word'],
            ko_score_completeness=ko_result['completeness'],
            ko_details=ko_result['details'],
            
            zh_score_total=zh_result['total'],
            zh_score_length=zh_result['length'],
            zh_score_chinese_ratio=zh_result['chinese_ratio'],
            zh_score_noun=zh_result['noun'],
            zh_score_predicate=zh_result['predicate'],
            zh_score_content_word=zh_result['content_word'],
            zh_score_completeness=zh_result['completeness'],
            zh_details=zh_result['details'],
            
            align_score_total=align_result['total'],
            align_score_length=align_result['length_match'],
            align_score_special=align_result['special_elements'],
            align_score_punctuation=align_result['punctuation_match'],
            align_score_question_mark=align_result['question_mark_match'],
            align_score_sentence_type=align_result['sentence_type_match'],
            align_score_bilingual=align_result['bilingual_note'],
            align_details=align_result['details'],
            
            final_score=final_score,
            final_details=final_details
        )


# ===========================================================================
# Excel 处理模块
# ===========================================================================

def process_excel(input_path: str, output_path: str = None, log_callback=None, ko_column=None, zh_column=None):
    """处理Excel文件"""
    if not HAS_OPENPYXL:
        print("错误: openpyxl 未安装")
        return
    
    input_path = Path(input_path)
    if not input_path.exists():
        print(f"错误: 文件不存在: {input_path}")
        return
    
    log_msg = f"读取文件: {input_path}"
    print(log_msg)
    if log_callback:
        log_callback(log_msg)
    
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active

    ko_col, ko_source, ko_auto = find_column_by_headers(ws, KO_COLUMN_CANDIDATES, ko_column or DEFAULT_KO_COLUMN)
    zh_col, zh_source, zh_auto = find_column_by_headers(ws, ZH_COLUMN_CANDIDATES, zh_column or DEFAULT_ZH_COLUMN)

    log_msg = (
        f"列匹配结果: 韩文列={ko_col}({ko_source})"
        f"{' 自动匹配' if ko_auto else ' 固定回退'} | "
        f"中文列={zh_col}({zh_source})"
        f"{' 自动匹配' if zh_auto else ' 固定回退'}"
    )
    print(log_msg)
    if log_callback:
        log_callback(log_msg)
    
    max_col = ws.max_column
    
    new_headers = [
        "韩文总分", "韩文长度", "韩文比例", "韩文名词", "韩文动形", "韩文内容词", "韩文完整",
        "中文总分", "中文长度", "中文比例", "中文名词", "中文谓词", "中文内容词", "中文完整",
        "对齐总分", "对齐长度", "对齐特殊元素", "对齐标点", "对齐问号", "对齐句型", "对齐双语",
        "最终总分", "评分详情"
    ]
    total_column_offsets = [0, 7, 14, 21]
    total_fill = PatternFill(fill_type="solid", fgColor="FFF200")
    total_font = Font(name="Microsoft YaHei", bold=True)
    total_alignment = Alignment(horizontal="center", vertical="center")
    
    start_col = max_col + 1
    for i, header in enumerate(new_headers):
        header_cell = ws.cell(1, start_col + i)
        header_cell.value = header
        if i in total_column_offsets:
            header_cell.fill = total_fill
            header_cell.font = total_font
            header_cell.alignment = total_alignment
    
    scorer = ParallelCorpusScorer()
    
    total_rows = ws.max_row - 1
    log_msg = f"开始处理 {total_rows} 行数据..."
    print(log_msg)
    if log_callback:
        log_callback(log_msg)
    
    # 用于统计平均分
    ko_scores = []
    zh_scores = []
    align_scores = []
    final_scores = []
    
    for row in range(2, ws.max_row + 1):
        if row % 100 == 0:
            log_msg = f"  处理到第 {row} 行..."
            print(log_msg)
            if log_callback:
                log_callback(log_msg)
        
        ko_text = ws.cell(row, ko_col).value or ""
        zh_text = ws.cell(row, zh_col).value or ""
        
        result = scorer.score(str(ko_text), str(zh_text))
        
        # 收集分数用于统计
        ko_scores.append(result.ko_score_total)
        zh_scores.append(result.zh_score_total)
        align_scores.append(result.align_score_total)
        final_scores.append(result.final_score)
        
        numeric_values = [
            result.ko_score_total, result.ko_score_length, result.ko_score_hangul,
            result.ko_score_noun, result.ko_score_verb_adj, result.ko_score_content_word,
            result.ko_score_completeness,
            result.zh_score_total, result.zh_score_length, result.zh_score_chinese_ratio,
            result.zh_score_noun, result.zh_score_predicate, result.zh_score_content_word,
            result.zh_score_completeness,
            result.align_score_total, result.align_score_length, result.align_score_special,
            result.align_score_punctuation, result.align_score_question_mark,
            result.align_score_sentence_type, result.align_score_bilingual,
            result.final_score
        ]
        for idx, value in enumerate(numeric_values):
            cell = ws.cell(row, start_col + idx)
            cell.value = value
            cell.number_format = '0.00'
            if idx in total_column_offsets:
                cell.fill = total_fill
                cell.font = total_font
                cell.alignment = total_alignment
        ws.cell(row, start_col + len(numeric_values)).value = result.final_details
    
    avg_ko = avg_zh = avg_align = avg_final = None
    if ko_scores:
        avg_ko = sum(ko_scores) / len(ko_scores)
        avg_zh = sum(zh_scores) / len(zh_scores)
        avg_align = sum(align_scores) / len(align_scores)
        avg_final = sum(final_scores) / len(final_scores)

        if "summary" in wb.sheetnames:
            wb.remove(wb["summary"])
        summary_ws = wb.create_sheet("summary", 0)
        summary_items = [
            ("韩文平均分", avg_ko),
            ("中文平均分", avg_zh),
            ("对齐平均分", avg_align),
            ("平行语料平均分", avg_final),
        ]
        title_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        item_fill = PatternFill(fill_type="solid", fgColor="F7FBFF")
        thin_side = Side(style="thin", color="B7C9D6")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        summary_ws.merge_cells("A1:B1")
        title_cell = summary_ws["A1"]
        title_cell.value = "评分统计信息"
        title_cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=14)
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.border = thin_border
        summary_ws.row_dimensions[1].height = 24

        header_cells = [summary_ws["A3"], summary_ws["B3"]]
        header_values = ["项目", "数值"]
        for cell, value in zip(header_cells, header_values):
            cell.value = value
            cell.font = Font(name="Microsoft YaHei", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for idx, (label, value) in enumerate(summary_items, start=4):
            item_cell = summary_ws.cell(idx, 1)
            value_cell = summary_ws.cell(idx, 2)
            item_cell.value = label
            value_cell.value = value
            value_cell.number_format = "0.00"

            item_cell.font = Font(name="Microsoft YaHei")
            value_cell.font = Font(name="Microsoft YaHei")
            item_cell.fill = item_fill
            value_cell.fill = item_fill
            item_cell.alignment = Alignment(horizontal="left", vertical="center")
            value_cell.alignment = Alignment(horizontal="right", vertical="center")
            item_cell.border = thin_border
            value_cell.border = thin_border

        summary_ws.column_dimensions["A"].width = 18
        summary_ws.column_dimensions["B"].width = 12
        summary_ws.freeze_panes = "A3"
        summary_ws.sheet_view.showGridLines = False
        summary_ws.title = "summary"

    if output_path is None:
        output_path = input_path.parent / f"Scored_{input_path.name}"
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    
    log_msg = f"✅ 完成！结果已保存到: {output_path}"
    print(log_msg)
    if log_callback:
        log_callback(log_msg)
    
    if avg_ko is not None:
        log_msg = "\n" + "=" * 60
        log_msg += "\n📊 评分统计信息"
        log_msg += "\n" + "=" * 60
        log_msg += f"\n韩文平均分: {avg_ko:.2f}"
        log_msg += f"\n中文平均分: {avg_zh:.2f}"
        log_msg += f"\n对齐平均分: {avg_align:.2f}"
        log_msg += f"\n平行语料平均分: {avg_final:.2f}"
        log_msg += "\n" + "=" * 60
        
        print(log_msg)
        if log_callback:
            log_callback(log_msg)
        
        return avg_ko, avg_zh, avg_align, avg_final
    return None


# ===========================================================================
# GUI 界面模块
# ===========================================================================

class ScorerGUI:
    """评分程序图形界面"""
    
    def __init__(self, root):
        self.root = root
        self.root.title("韩中平行语料库质量评分程序")
        self.root.geometry("700x550")
        
        self.input_file = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.output_filename = tk.StringVar()
        
        self.scorer = ParallelCorpusScorer()
        self.is_running = False
        
        self.setup_ui()
    
    def setup_ui(self):
        """设置界面"""
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        title_label = ttk.Label(
            main_frame,
            text="韩中平行语料库质量评分程序",
            font=("Microsoft YaHei", 16, "bold")
        )
        title_label.grid(row=0, column=0, columnspan=3, pady=(0, 20))
        
        ttk.Label(main_frame, text="1. 选择输入Excel文件:", font=("Microsoft YaHei", 10, "bold")).grid(
            row=1, column=0, sticky=tk.W, pady=(0, 5)
        )
        
        input_frame = ttk.Frame(main_frame)
        input_frame.grid(row=2, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 15))
        
        ttk.Entry(input_frame, textvariable=self.input_file, width=50).grid(
            row=0, column=0, padx=(0, 5)
        )
        ttk.Button(input_frame, text="浏览...", command=self.browse_input).grid(
            row=0, column=1
        )
        
        ttk.Label(main_frame, text="2. 选择输出文件夹:", font=("Microsoft YaHei", 10, "bold")).grid(
            row=3, column=0, sticky=tk.W, pady=(0, 5)
        )
        
        output_dir_frame = ttk.Frame(main_frame)
        output_dir_frame.grid(row=4, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 15))
        
        ttk.Entry(output_dir_frame, textvariable=self.output_dir, width=50).grid(
            row=0, column=0, padx=(0, 5)
        )
        ttk.Button(output_dir_frame, text="浏览...", command=self.browse_output_dir).grid(
            row=0, column=1
        )
        
        ttk.Label(main_frame, text="3. 输出文件名 (自动添加 Scored_ 前缀):", font=("Microsoft YaHei", 10, "bold")).grid(
            row=5, column=0, sticky=tk.W, pady=(0, 5)
        )
        
        ttk.Entry(main_frame, textvariable=self.output_filename, width=50).grid(
            row=6, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(0, 20)
        )
        
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=7, column=0, columnspan=3, pady=(0, 15))
        
        self.start_button = ttk.Button(
            button_frame,
            text="开始评分",
            command=self.start_scoring,
            width=15
        )
        self.start_button.grid(row=0, column=0, padx=(0, 10))
        
        ttk.Button(
            button_frame,
            text="退出",
            command=self.root.quit,
            width=15
        ).grid(row=0, column=1)
        
        ttk.Label(main_frame, text="处理进度:").grid(row=8, column=0, sticky=tk.W, pady=(0, 5))
        self.progress = ttk.Progressbar(main_frame, mode='determinate', length=500)
        self.progress.grid(row=9, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 10))
        
        ttk.Label(main_frame, text="处理日志:").grid(row=10, column=0, sticky=tk.W, pady=(0, 5))
        self.log_text = scrolledtext.ScrolledText(
            main_frame,
            width=60,
            height=10,
            wrap=tk.WORD
        )
        self.log_text.grid(row=11, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(11, weight=1)
    
    def log(self, message):
        """添加日志"""
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.root.update_idletasks()
    
    def browse_input(self):
        """浏览输入文件"""
        filename = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if filename:
            self.input_file.set(filename)
            
            input_path = Path(filename)
            self.output_dir.set(str(input_path.parent))
            
            if not self.output_filename.get():
                self.output_filename.set(input_path.name)
    
    def browse_output_dir(self):
        """浏览输出文件夹"""
        dirname = filedialog.askdirectory(title="选择输出文件夹")
        if dirname:
            self.output_dir.set(dirname)
    
    def start_scoring(self):
        """开始评分"""
        if self.is_running:
            return
        
        input_path = self.input_file.get()
        output_dir = self.output_dir.get()
        output_name = self.output_filename.get()
        
        if not input_path:
            messagebox.showerror("错误", "请选择输入文件！")
            return
        
        if not output_dir:
            messagebox.showerror("错误", "请选择输出文件夹！")
            return
        
        input_path_obj = Path(input_path)
        if not input_path_obj.exists():
            messagebox.showerror("错误", "输入文件不存在！")
            return
        
        if not output_name:
            output_name = input_path_obj.name
        
        output_name_obj = Path(output_name)
        final_output_name = f"Scored_{output_name_obj.stem}{output_name_obj.suffix}"
        output_path = Path(output_dir) / final_output_name
        
        self.is_running = True
        self.start_button.config(state="disabled")
        self.log_text.delete(1.0, tk.END)
        
        thread = threading.Thread(
            target=self.run_scoring,
            args=(str(input_path_obj), str(output_path))
        )
        thread.daemon = True
        thread.start()
    
    def run_scoring(self, input_path: str, output_path: str):
        """运行评分"""
        try:
            self.log("=" * 60)
            self.log("开始处理...")
            self.log(f"输入文件: {input_path}")
            self.log(f"输出文件: {output_path}")
            
            if not HAS_OPENPYXL:
                self.log("错误: openpyxl 未安装")
                messagebox.showerror("错误", "openpyxl 未安装！")
                return
            
            # 使用process_excel函数处理
            result = process_excel(input_path, output_path, self.log)
            
            if result:
                avg_ko, avg_zh, avg_align, avg_final = result
                # 显示统计信息弹窗
                stats_msg = f"📊 评分统计信息\n\n"
                stats_msg += f"韩文平均分: {avg_ko:.2f}\n"
                stats_msg += f"中文平均分: {avg_zh:.2f}\n"
                stats_msg += f"对齐平均分: {avg_align:.2f}\n"
                stats_msg += f"平行语料平均分: {avg_final:.2f}"
                
                self.log("=" * 60)
                self.log("处理完成！")
                self.log("=" * 60)
                
                messagebox.showinfo("评分完成", f"处理完成！\n\n结果已保存到:\n{output_path}\n\n{stats_msg}")
            else:
                self.log("=" * 60)
                self.log("处理完成！")
                self.log("=" * 60)
                messagebox.showinfo("完成", f"处理完成！\n结果已保存到:\n{output_path}")
            
        except Exception as e:
            self.log(f"❌ 错误: {str(e)}")
            messagebox.showerror("错误", f"处理出错:\n{str(e)}")
        
        finally:
            self.is_running = False
            self.start_button.config(state="normal")


# ===========================================================================
# 主程序
# ===========================================================================

def main():
    import sys
    
    if len(sys.argv) > 1 and sys.argv[1] != "--gui":
        print("=" * 60)
        print("韩中平行语料库质量评分程序")
        print("=" * 60)
        
        default_input = r"d:\PythonProject\corpus_trusted\ko_corpus_10000_20260615_232407_translated.xlsx"
        
        input_path = sys.argv[1] if len(sys.argv) > 1 else default_input
        output_path = sys.argv[2] if len(sys.argv) > 2 else None
        
        process_excel(input_path, output_path)
    else:
        if not HAS_GUI:
            print("错误: tkinter 不可用")
            return
        
        root = tk.Tk()
        app = ScorerGUI(root)
        
        default_input = r"d:\PythonProject\corpus_trusted\ko_corpus_10000_20260615_232407_translated.xlsx"
        if Path(default_input).exists():
            app.input_file.set(default_input)
            app.output_dir.set(str(Path(default_input).parent))
            app.output_filename.set(Path(default_input).name)
        
        root.mainloop()


if __name__ == "__main__":
    main()
